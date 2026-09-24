from io import BytesIO

from fastapi import HTTPException
import pymupdf
from docx import Document


MAX_FILE_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}


def _get_extension(filename: str) -> str:
    if not filename or "." not in filename:
        return ""

    return filename.rsplit(".", 1)[-1].lower().strip()


def _extract_pdf(data: bytes) -> str:
    # A normal PDF file starts with the %PDF signature.
    if not data.startswith(b"%PDF"):
        raise HTTPException(
            status_code=400,
            detail="The uploaded file does not appear to be a valid PDF.",
        )

    document = pymupdf.open(
        stream=data,
        filetype="pdf",
    )

    chunks = []

    try:
        for page in document:
            page_text = page.get_text("text").strip()

            if page_text:
                chunks.append(page_text)

    finally:
        document.close()

    return "\n".join(chunks).strip()


def _extract_docx(data: bytes) -> str:
    # DOCX files are ZIP-based and normally start with the PK signature.
    if not data.startswith(b"PK"):
        raise HTTPException(
            status_code=400,
            detail="The uploaded file does not appear to be a valid DOCX file.",
        )

    document = Document(BytesIO(data))

    chunks = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            chunks.append(paragraph_text)

    # Extract text stored inside tables.
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                chunks.append(" | ".join(cells))

    return "\n".join(chunks).strip()


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig").strip()

    except UnicodeDecodeError:
        raise HTTPException(
            status_code=400,
            detail="The TXT file must use UTF-8 text encoding.",
        )


def extract_text_from_file(
    data: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:

    if not data:
        raise HTTPException(
            status_code=400,
            detail="No file uploaded.",
        )

    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File is too large. Maximum size is 10 MB.",
        )

    extension = _get_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Please upload PDF, DOCX, or TXT.",
        )

    try:
        if extension == "pdf":
            text = _extract_pdf(data)

        elif extension == "docx":
            text = _extract_docx(data)

        else:
            text = _extract_txt(data)

    except HTTPException:
        raise

    except Exception as exc:
        # Log the technical error on the backend without exposing
        # implementation details to the browser.
        print(f"Document extraction error: {exc}")

        raise HTTPException(
            status_code=400,
            detail="Could not read the uploaded file.",
        ) from exc

    if not text:
        raise HTTPException(
            status_code=422,
            detail="The uploaded file contains no extractable text.",
        )

    return text